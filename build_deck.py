from pathlib import Path
import math, json, io, re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties, fontManager
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement

BASE=Path(__file__).resolve().parent
ROOT=BASE/'assets'
ROOT.mkdir(exist_ok=True)
OUT=BASE
OUT.mkdir(exist_ok=True)
BLUE='587C95'; ORANGE='DA501B'; INK='243746'; MUTED='607281'
PALE='EDF3F6'; WARM='FCF0E9'; GREEN='3E8875'; WHITE='FFFFFF'; GRID='D9E3E9'
FONT='Noto Sans CJK SC'
font_file='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
fontManager.addfont(font_file)
PLOT_FONT=FontProperties(fname=font_file).get_name()
plt.rcParams.update({'font.family':PLOT_FONT,'axes.unicode_minus':False,'font.size':12,
                     'axes.spines.top':False,'axes.spines.right':False})
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Catch-All or No-False-Alarm：不完美审计下的社会责任供应链管理'
prs.core_properties.subject='中文组会汇报：研究背景、研究问题、模型构建、最优策略、拓展'
prs.core_properties.author='组会汇报'
notes_all=[]; figures_used=set(); math_count=0

def C(hex): return RGBColor.from_string(hex)

def tx(s,x,y,w,h,text,size=20,color=INK,bold=False,align=PP_ALIGN.LEFT,margin=0.02):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin);tf.margin_top=tf.margin_bottom=Inches(margin)
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align
        p.space_after=Pt(3);p.line_spacing=1.05
        r=p.add_run();r.text=line;r.font.name=FONT;r.font.size=Pt(size)
        r.font.bold=bold;r.font.color.rgb=C(color)
        rp=r._r.get_or_add_rPr();ea=OxmlElement('a:ea');ea.set('typeface',FONT);rp.append(ea)
    return sh

def box(s,x,y,w,h,fill=PALE,line=None,radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                          Inches(x),Inches(y),Inches(w),Inches(h))
    if radius:
        sh.adjustments[0]=0.08
    sh.fill.solid();sh.fill.fore_color.rgb=C(fill)
    if line: sh.line.color.rgb=C(line);sh.line.width=Pt(1)
    else: sh.line.fill.background()
    return sh

def line(s,x1,y1,x2,y2,color=GRID,width=1):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=C(color);sh.line.width=Pt(width);return sh

def arrow(s,x,y,w=.42,h=.26,color=BLUE):
    sh=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=C(color);sh.line.fill.background();return sh

def pic(s,path,x,y,w,h):
    path=Path(path)
    with Image.open(path) as im: ar=im.width/im.height
    pw=min(w,h*ar);ph=pw/ar
    return s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))

def eq(s,latex,x,y,w,h,color=INK):
    global math_count
    math_count+=1;path=ROOT/f'math_{math_count:03}.png'
    f=plt.figure(figsize=(12,1.2),dpi=250)
    f.text(.01,.5,'$'+latex+'$',fontsize=28,color='#'+color,va='center')
    f.savefig(path,dpi=250,bbox_inches='tight',pad_inches=.05,transparent=True);plt.close(f)
    return pic(s,path,x,y,w,h)

def strip(s,text,y=6.5,x=.68,w=11.97,size=18,color=BLUE):
    y=min(y,6.52)
    box(s,x,y,w,.49,fill=PALE if color==BLUE else WARM)
    tx(s,x+.16,y+.035,w-.32,.445,text,size=size,color=color,bold=True)

def card(s,x,y,w,h,label,body,color=BLUE,size=19):
    box(s,x,y,w,h,fill=PALE if color==BLUE else WARM,radius=True)
    box(s,x,y,.055,h,fill=color)
    short=h<=1.45
    label_size=18 if short else 20
    body_y=.55 if short else .66
    body_size=min(size,16) if short else size
    available=h-body_y-.08
    def estimate_required(fs):
        capacity=(w-.4)*72/fs
        lines=sum(max(1,math.ceil(sum(1 if ord(c)>255 else .52 for c in p)/capacity)) for p in body.split('\n'))
        return lines*fs*1.45/72+(len(body.split('\n'))-1)*3/72
    while body_size>14 and estimate_required(body_size)>available:
        body_size-=.5
    tx(s,x+.2,y+.10,w-.38,.43,label,size=label_size,color=color,bold=True)
    tx(s,x+.2,y+body_y,w-.4,available,body,size=body_size)

def bullets(s,items,x=.75,y=2.63,w=11.8,gap=.9,size=21):
    for i,item in enumerate(items):
        box(s,x,y+i*gap+.12,.095,.095,fill=ORANGE)
        tx(s,x+.26,y+i*gap,w-.26,gap-.08,item,size=size)

def table(s,headers,rows,x=.7,y=2.55,w=11.93,col=None,rowh=.69,size=17):
    cols=len(headers); widths=col or [w/cols]*cols
    sh=s.shapes.add_table(len(rows)+1,cols,Inches(x),Inches(y),Inches(w),Inches(rowh*(len(rows)+1)))
    t=sh.table
    for j,wd in enumerate(widths):t.columns[j].width=Inches(wd)
    for i,row in enumerate([headers]+rows):
        t.rows[i].height=Inches(rowh)
        for j,txt in enumerate(row):
            c=t.cell(i,j);c.margin_left=c.margin_right=Inches(.13);c.margin_top=c.margin_bottom=Inches(.08)
            c.fill.solid();c.fill.fore_color.rgb=C(BLUE if i==0 else (PALE if i%2 else WHITE))
            c.vertical_anchor=MSO_ANCHOR.MIDDLE;c.text=str(txt)
            for p in c.text_frame.paragraphs:
                p.space_after=Pt(1);p.line_spacing=1.08
                for r in p.runs:
                    r.font.name=FONT;r.font.size=Pt(size if i else size+1);r.font.bold=(i==0);r.font.color.rgb=C(WHITE if i==0 else INK)
                    rp=r._r.get_or_add_rPr();ea=OxmlElement('a:ea');ea.set('typeface',FONT);rp.append(ea)
    return sh

def slide(section,title,source='',notes='',appendix=False):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(ROOT/'image2.png'),0,0,width=prs.slide_width,height=prs.slide_height)
    tx(s,1.63,.64,8.8,.55,section,27,bold=True)
    tx(s,9.45,.84,3.14,.24,'文献精读  /  POM 2026',10.5,color=MUTED,align=PP_ALIGN.RIGHT)
    if title:tx(s,.68,1.87,12.0,.55,title,24,bold=True)
    line(s,.68,7.08,12.67,7.08)
    tx(s,.72,7.13,11.0,.23,source or 'Li & Li (2026) · DOI: 10.1177/10591478261468995',10,color=MUTED)
    tx(s,11.65,7.11,1.0,.28,f'{len(prs.slides):02d}',11,color=BLUE,align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text=notes
    notes_all.append({'slide':len(prs.slides),'section':section,'title':title,'notes':notes,'appendix':appendix})
    return s

def crop_panel(src,name,rect):
    # Crops identify subpanels of the supplied scholarly figure; no data are altered.
    with Image.open(ROOT/src) as im: im.crop(rect).save(ROOT/name)

# Original figures: keep complete versions and enlarge panels where useful.
with Image.open(ROOT/'fig2.png') as im:
    W,H=im.size
    im.crop((0,0,int(W*.498),int(H*.89))).save(ROOT/'fig2a.png')
    im.crop((int(W*.499),0,W,int(H*.89))).save(ROOT/'fig2b.png')
with Image.open(ROOT/'fig3.png') as im:
    W,H=im.size
    im.crop((0,0,int(W*.499),int(H*.88))).save(ROOT/'fig3a.png')

# Independently calculated illustrations, explicitly distinguished from paper figures.
e=.5;Da=3.;Db=5.5;ds=1.6;dr=2.
bbar=(Db-(2-e)*dr-e*ds)/(Db-dr-ds)
def loss_at(b):
    terms=np.array([(b-e)*dr+(1-b)*Db,(b-e)*ds+(1-b)*Db,
                    (2-e-b)*dr,(2-e-b)*ds+(1-b)*Da])
    j=int(terms.argmin());return float(terms[j]), ['NFA+r','NFA+s','CA+r','CA+s'][j]
bs=np.linspace(.50001,1,501)
fig,axs=plt.subplots(1,2,figsize=(11,3.6),layout='constrained')
axs[0].plot(bs,(2-e-bs)*dr,color='#'+BLUE,label='Catch-all + 整改',lw=2)
axs[0].plot(bs,(1-e)*ds+(1-bs)*(Db-ds),color='#'+ORANGE,label='No-false-alarm + 替换',lw=2)
axs[0].axvline(bbar,color='#'+MUTED,ls='--',lw=1);axs[0].set(xlabel=r'可达准确度 $b_0$',ylabel='企业期望损失',title='最优策略在两条损失曲线之间切换')
axs[0].legend(frameon=False,fontsize=10)
axs[1].plot(bs[bs<=bbar],np.zeros(sum(bs<=bbar)),color='#'+BLUE,lw=2.5)
axs[1].plot(bs[bs>bbar],1-bs[bs>bbar],color='#'+ORANGE,lw=2.5)
axs[1].plot([bbar,bbar],[0,1-bbar],color='#'+MUTED,ls='--')
axs[1].scatter([bbar],[1-bbar],color='#'+ORANGE,s=22)
axs[1].set(xlabel=r'可达准确度 $b_0$',ylabel='违规产品进入市场的概率 L*',title='切换点可能出现向上跳跃',ylim=(-.01,.18))
for ax in axs:ax.grid(alpha=.15)
fig.savefig(ROOT/'benchmark_calc.png',dpi=240,bbox_inches='tight');plt.close(fig)

# Both suppliers audited: solve all affine policy pairs along the feasible circle.
budgets=np.linspace(.25001,1,700)
alloc=[]
for B in budgets:
    b1=np.linspace(math.sqrt(B),min(1,math.sqrt(2*B-e*e)),8001)
    b2=np.sqrt(np.maximum(0,2*B-b1*b1))
    def vals(v):
        return np.minimum.reduce([(v-e)*dr+(1-v)*Db,(v-e)*ds+(1-v)*Db,
                                 (2-e-v)*dr,(2-e-v)*ds+(1-v)*Da])
    z=vals(b1)+vals(b2);k=int(np.argmin(z));v1=float(b1[k]);v2=float(b2[k])
    q1=0 if loss_at(v1)[1].startswith('CA') else 1-v1
    q2=0 if loss_at(v2)[1].startswith('CA') else 1-v2
    alloc.append((B,v1,v2,1-(1-q1)*(1-q2)))
data=np.array(alloc)
fig,ax=plt.subplots(figsize=(8.0,3.4),layout='constrained')
ax.plot(data[:,0],data[:,1],label=r'供应商 1：$b_1^*$',color='#'+ORANGE,lw=2.4)
ax.plot(data[:,0],data[:,2],label=r'供应商 2：$b_2^*$',color='#'+BLUE,lw=2.4)
ax.plot(data[:,0],np.sqrt(data[:,0]),ls=':',color='#'+MUTED,lw=1.4,label='平均分配基准 √B')
ax.set(xlabel='总预算 B（条件：审计两家）',ylabel='最优审计准确度',ylim=(.48,1.025))
ax.legend(frameon=False,fontsize=10,loc='lower right');ax.grid(alpha=.15)
fig.savefig(ROOT/'allocation_calc.png',dpi=240,bbox_inches='tight');plt.close(fig)
fig,ax=plt.subplots(figsize=(8,3.4),layout='constrained')
ax.plot(data[:,0],data[:,3],color='#'+ORANGE,lw=2.4)
ax.set(xlabel='总预算 B（条件：审计两家）',ylabel='违规最终产品进入市场概率 L*',ylim=(-.014,.12))
ax.grid(alpha=.15)
fig.savefig(ROOT/'social_calc.png',dpi=240,bbox_inches='tight');plt.close(fig)

# 01 Cover
s=slide('组会文献汇报','',notes='建议用时 30 秒。论文标题中的 Catch-all 指优先消除漏报，No-false-alarm 指优先消除误报。不要把无误报解释为完美审计。封面汇报人与日期为可编辑字段，避免沿用模板的旧汇报日期。')
tx(s,1.03,2.04,11.28,1.53,'Catch-All or No-False-Alarm?\nManaging Socially Responsible Supply Chains with Imperfect Audits',25,bold=True,align=PP_ALIGN.CENTER)
tx(s,1.15,3.83,11.03,1.24,'全面筛查还是避免误报？\n不完美审计下的社会责任供应链管理',29,bold=True,color=BLUE,align=PP_ALIGN.CENTER)
tx(s,1.65,5.27,10.04,.87,'Chengzhang Li · Yuanchen Li\nProduction and Operations Management，2026',18,align=PP_ALIGN.CENTER)
tx(s,2.7,6.36,7.9,.52,'汇报人：待填写        汇报日期：待填写',18,align=PP_ALIGN.CENTER)

# 02 Article information
s=slide('论文信息','文章定位：从“只防漏报”走向“两类错误联合管理”','论文首页与摘要，pp. 1–3',
        '建议用时 30 秒。作者为 Chengzhang Li 和 Yuanchen Li，分别来自上海交通大学和同济大学相关单位。本文 2026 年发表于 POM，所提供 PDF 是 21 页的在线文章版本。左边保留原文首页片段。基准模型中只有买方做序贯决策，供应商的合规概率外生；供应商策略性努力在第 6.1 节引入。')
pic(s,ROOT/'paper_header.png',.75,2.55,6.7,3.95)
card(s,7.7,2.6,4.75,1.3,'研究对象','一个买方 + 两个投入品供应商',size=18)
card(s,7.7,4.06,4.75,2.33,'核心贡献','将误报、漏报与事后治理联立，解释预算分配、审计覆盖与社会效果的反直觉变化。',color=ORANGE,size=19)

# 03 Roadmap
s=slide('目录','沿用模板的七部分结构，突出模型与分析','汇报安排：正文约 30 分钟；附录供讨论查阅',
        '建议用时 20 秒。七部分与用户提供的模板一致。正文 32 页，其中模型构建与均衡分析是重点。附录 6 页放完整条件、计算依据和复杂拓展命题，不计入常规汇报时间。可根据现场讨论调整每页速度。')
sections=[('一、研究背景','2 分钟'),('二、研究问题','约 1 分钟'),('三、文献综述','约 1 分钟'),('四、模型构建','8–9 分钟'),('五、均衡分析','11–12 分钟'),('六、扩展模型','约 4 分钟'),('七、结论与未来展望','2 分钟')]
for i,(a,b) in enumerate(sections):
    col=0 if i<4 else 1;row=i if col==0 else i-4
    x=.92+6.25*col;y=2.72+row*.9
    tx(s,x,y,4.55,.53,a,25,color=BLUE if i==3 else INK,bold=(i in [3,4]))
    tx(s,x+4.15,y+.08,1.43,.46,b,16.5,color=MUTED)

# 04 Background
s=slide('一、研究背景','供应商违规会传导到品牌；审计之后还必须决定如何治理','引言，pp. 1–2；案例按论文叙述整理',
        '建议用时 50 秒。全球分工使品牌大量采购海外投入品，但供应商的劳动、人权、安全和环境问题可转化为买方声誉与经济损失。论文列举 Apple、IKEA、Walmart、Starbucks 的责任审计，以及整改和替换两类后续行动。案例仅用于说明动机，不是本文模型的实证检验。对 Dyson 的表述保持“被指控”，不把指控当作已判定事实。')
card(s,.7,2.65,3.86,2.95,'风险传导','上游违规\n→ 媒体 / 监管发现\n→ 下游品牌受损\n例：Dyson 供应链被指控存在强迫劳动。',size=19)
card(s,4.74,2.65,3.86,2.95,'常规应对','Apple、IKEA 等开展责任审计。\n审计未通过后：\n整改原供应商，或改用合规供应商。',size=19)
card(s,8.78,2.65,3.86,2.95,'现实约束','审计不完美，而且预算有限。\n既可能“放过违规者”，也可能“误判合规者”。\n例：Maggi 面条检测争议提示误报风险。',color=ORANGE,size=19)
strip(s,'研究动机：审计范围、错误控制与事后治理，应当作为一套方案设计。',y=6.23,size=20)

# 05 errors
s=slide('一、研究背景','误报与漏报对应不同损失；“审计未通过”不等于真实违规','两类错误定义与损失机制，pp. 2、5–6',
        '建议用时 60 秒。表中是给定真实类型后的条件概率，α 和 β 不是无条件出错概率。Type I 是合规被判失败，Type II 是违规被判通过。误报后无论整改还是替换，都可能浪费治理成本；但本文的一类声誉损失 Dα 只在错误替换发生并被外部观察后产生。整改保持关系，在模型中避免该项损失。论文以 Maggi 检测争议和供应商指控争议说明误报风险。')
table(s,['真实状态 / 审计结果','通过 o = 0','未通过 o = 1'],[
 ['合规（先验概率 e）','正确通过：1 − αᵢ','误报 / Type I：αᵢ'],
 ['违规（先验概率 1 − e）','漏报 / Type II：βᵢ','正确发现：1 − βᵢ']],rowh=.76,size=20,col=[4.7,3.62,3.61])
card(s,.72,5.03,5.75,1.25,'误报的后果','不必要治理；错误替换还可能触发 Dα。',size=17)
card(s,6.68,5.03,5.94,1.25,'漏报的后果','违规投入品留在供应链；曝光后损失 Dβ。',color=ORANGE,size=17)
tx(s,.91,6.53,11.55,.35,'Dα 与 Dβ 已包含外部发现概率；不是把每次误判都视为已被公众发现。',17,color=MUTED)

# 06 RQ
s=slide('二、研究问题','有限预算下，买方如何联合设计审计与治理？','研究问题，pp. 2–3',
        '建议用时 50 秒。先提出四个问题而不是立即给结论。第一是控制哪一种错误，第二是失败后整改还是替换，第三是在供应商间如何分配资源且审计多少家，第四是私益与社会效果是否一致。后续四组分析按这个顺序回答。')
for i,(a,b) in enumerate([
 ('错误控制','应优先消除误报，还是优先消除漏报？'),
 ('治理动作','审计未通过后，应整改原供应商还是替换它？'),
 ('资源与范围','如何分配预算？是否必须平均分配、覆盖更多供应商？'),
 ('社会效果','预算增加、企业收益改善，是否一定减少市场中的违规产品？')]):
    y=2.61+i*.93
    box(s,.74,y,.63,.6,fill=BLUE if i<3 else ORANGE)
    tx(s,.74,y+.07,.63,.4,f'{i+1:02}',22,color=WHITE,bold=True,align=PP_ALIGN.CENTER)
    tx(s,1.58,y+.03,2.0,.42,a,22,color=BLUE,bold=True)
    tx(s,3.6,y+.03,8.79,.66,b,21)

# 07 literature
s=slide('三、文献综述','研究缺口：错误率、治理动作与网络预算不能分开看','文献综述，pp. 3–4；此表为汇报整理',
        '建议用时 50 秒。责任采购文献研究审计激励、信息披露和审计规避；响应治理文献研究发现问题后如何整改或退出；网络审计文献研究协作、结构和优先级。Guo 等 2017 已考虑两类错误，但错误率外生，重点在认证信号。本文的定位是内生选择错误组合，同时选择治理方案及网络资源分配。不要声称本文首次提出两类审计错误。')
table(s,['文献线索','代表研究（按原文）','本文补充'],[
 ['责任采购与审计','Plambeck & Taylor (2016)\nCho et al. (2019)','从主要关注漏报，转向误报—漏报权衡'],
 ['审计后的响应治理','Kalkanci & Plambeck (2020a)\nZhang et al. (2022)','将整改 / 替换与错误率联合优化'],
 ['网络审计与资源配置','Caro et al. (2018)\nChen et al. (2020a)','装配网络中研究预算分配与覆盖数'],
 ['不准确认证与分类','Guo et al. (2017)\nCezar et al. (2020)','错误率内生；连接治理成本与社会后果']],col=[3.05,4.5,4.38],rowh=.79,size=16)
strip(s,'文章定位：不完美审计 + 两类治理 + 有限预算 + 装配网络。',y=6.61,size=18)

# 08 network
s=slide('四、模型构建','装配网络：一个买方需要两个不同投入品','模型设定，§3，p. 5；网络示意为汇报重绘',
        '建议用时 60 秒。两个供应商不是互相替代的货源，各提供一个不同投入品，买方装配最终产品。基准假设供应商对称，合规状态由外生概率 e 决定，不主动调整努力。供应商状态独立的设定对应后文的乘积社会风险公式。整改和替换都能使被处理的投入品完全合规。主模型是单个买方的多阶段优化；只有内生责任努力拓展涉及买方与供应商的策略互动。')
card(s,.86,2.73,3.48,1.3,'供应商 1','投入品 1；合规概率 e',size=17)
card(s,.86,4.61,3.48,1.3,'供应商 2','投入品 2；合规概率 e',size=17)
arrow(s,4.67,3.28,.7,.32);arrow(s,4.67,5.16,.7,.32)
box(s,5.65,3.42,2.57,1.73,fill=BLUE,radius=True)
tx(s,5.76,3.64,2.35,1.22,'买方\n审计 + 装配',23,color=WHITE,bold=True,align=PP_ALIGN.CENTER)
arrow(s,8.47,4.15,.63,.32)
card(s,9.35,3.42,2.97,1.73,'最终产品','任一投入品违规，最终产品即不合规。',color=ORANGE,size=18)
strip(s,'基准：供应商不策略性调整努力；第 6.1 节再将合规努力内生化。',y=6.5,size=18)

# 09 notation
s=slide('四、模型构建','符号与决策变量：区分先验、错误率、治理成本和损失','符号整理自 §3，pp. 5–6',
        '建议用时 70 秒。x 是是否审计，α 和 β 是条件误判率，Ai 是整体正确分类概率。o 为结果，未审计时记为空。Δs 与 Δr 是买方的治理成本，Dα 和 Dβ 则是已乘外部发现概率的预期声誉或经济损失。审计预算 B 在基准模型中是约束，并不再次作为成本从收益中扣除；第 6.2 节才内生化预算并减去真实审计成本。')
table(s,['符号','含义','符号','含义'],[
 ['e ∈ [1/2, 1]','供应商先验合规概率','B','外生审计预算'],
 ['xᵢ ∈ {0, 1}','是否审计供应商 i','Aᵢ','整体审计准确度'],
 ['αᵢ / βᵢ','误报率 / 漏报率','oᵢ ∈ {0, 1, ∅}','通过 / 未通过 / 未审计'],
 ['fᵢ ∈ {s, r, ∅}','替换 / 整改 / 无行动','Δs / Δr','买方替换成本 / 整改成本'],
 ['Dα = rα D̃α','错误替换的预期损失','Dβ = rβ D̃β','遗漏违规的预期损失'],
 ['Π','无治理与损害时的交易利润','c','审计成本系数；基准归一化为 1']],rowh=.55,size=15.5,col=[2.25,3.78,2.18,3.72])
tx(s,.89,6.57,11.6,.37,'注意：αᵢ、βᵢ 是条件概率；eαᵢ、(1 − e)βᵢ 才是无条件误判概率。',17,color=BLUE,bold=True)

# 10 accuracy
s=slide('四、模型构建','准确度连接两类错误；提高准确度需要凸性的审计支出','审计技术与约束，§3，p. 5；公式对应原文',
        '建议用时 75 秒。Ai 等于两类真实状态下正确分类概率的加权和。固定 Ai 时，误报减少必然以漏报增加为代价，这是模型的等准确度线，不等于宣称任何真实审计技术都允许任意错误组合。二次支出使提升准确度的边际成本递增。原文叙述要求准确度优于先验，正式约束使用 ≥e；本汇报按正式公式列示，边界可视为弱不等式或极限。')
box(s,.77,2.63,7.11,1.16,fill=PALE)
eq(s,r'A_i=e(1-\alpha_i)+(1-e)(1-\beta_i)=1-e\alpha_i-(1-e)\beta_i',.93,2.88,6.79,.6)
box(s,.77,4.00,7.11,1.02,fill=WARM)
eq(s,r'e\alpha_i+(1-e)\beta_i=1-A_i',1.08,4.19,6.48,.61,color=ORANGE)
tx(s,.94,5.22,6.66,1.02,'固定准确度：降低误报 αᵢ，会提高漏报 βᵢ。\n增加资源：可同时降低两类错误。',21)
card(s,8.17,2.63,4.37,3.47,'成本与可行性','审计支出：cAᵢ² / 2\n基准：c = 1\n有效审计要求：Aᵢ ≥ e\n总支出不能超过 B。',size=20)
strip(s,'“高准确度”和“优先控制哪类错误”是两个不同的决策维度。',y=6.5,size=19)

# 11 timeline
s=slide('四、模型构建','决策时序：先设计审计，再观察结果，最后治理与交易','原文 Figure 1，p. 7；右侧为中文流程说明',
        '建议用时 70 秒。原图列出了真实状态、报告和后续损失的分支。按三阶段讲：买方先选审计对象和误差组合，看到审计结果后决定替换或整改，最终获取交易利润并承担外部发现带来的损失。买方不能直接观察真实状态，所以失败后决策依赖贝叶斯后验。替换或整改是审计结果之后的响应，而不是在结果实现前固定一个与结果无关的动作。')
pic(s,ROOT/'fig1.png',.72,2.60,8.11,3.97);figures_used.add(1)
for i,(a,b) in enumerate([('1. 审计阶段','选择 xᵢ、αᵢ、βᵢ\n满足预算与准确度约束'),('2. 治理阶段','观察 oᵢ\n未通过：选择替换或整改'),('3. 交易与损害','两个投入品装配并销售\n损害可能被外部发现')]):
    card(s,9.03,2.58+i*1.42,3.56,1.33,a,b,color=ORANGE if i==1 else BLUE,size=16)

# 12 events and branches
s=slide('四、模型构建','每个供应商的事件概率与后果：模型目标函数的来源','§3–4，pp. 5–7；表格为按原模型整理',
        '建议用时 70 秒。四个审计事件的无条件概率加起来等于 1。合规且失败的概率为 eα，若替换产生 Δs 及额外 Dα；若整改只有 Δr。违规且失败会被治理，避免违规投入品留在最终产品中。违规但通过的概率是 (1−e)β，只承担暴露时的预期 Dβ。未审计时不治理，期望损失为 (1−e)Dβ。表中的 D 都已加权了公众发现概率。')
table(s,['真实类型 × 审计结果','无条件概率','治理 / 预期后果'],[
 ['合规 × 通过','e(1 − αᵢ)','无治理；无错误损失'],
 ['合规 × 未通过','eαᵢ','替换：Δs + Dα；整改：Δr'],
 ['违规 × 通过','(1 − e)βᵢ','无治理；预期损失 Dβ'],
 ['违规 × 未通过','(1 − e)(1 − βᵢ)','替换：Δs；整改：Δr'],
 ['未审计','xᵢ = 0','无治理；预期损失 (1 − e)Dβ']],rowh=.66,size=17,col=[4.05,3.13,4.75])
strip(s,'“审计未通过”的总概率：pᵢ = eαᵢ + (1 − e)(1 − βᵢ)。',y=6.68,size=18)

# 13 posterior
s=slide('四、模型构建','贝叶斯更新决定治理：替换节省的钱，能否覆盖误判风险？','Lemma 1 / Proposition 3，pp. 7、10',
        '建议用时 95 秒。失败后可能是真违规，也可能是误报。后验合规概率等于误报事件的概率除以全部失败概率。替换的条件预期成本为 Δs 加后验合规概率乘 Dα，整改成本为 Δr。所以选择替换当且仅当替换成本优势超过错误替换的条件预期损失。若 pᵢ=0，失败事件不会发生，不需要定义该事件上的后验。装配网络下该阈值只取决于本供应商，治理动作可逐个求解。相等时买方无差异。')
eq(s,r'\Pr(C_i\mid o_i=1)=\frac{e\alpha_i}{e\alpha_i+(1-e)(1-\beta_i)}',.91,2.66,7.7,.95)
eq(s,r'\delta(\alpha_i,\beta_i)=\Pr(C_i\mid o_i=1)D_\alpha',.91,3.83,7.7,.7)
eq(s,r'\Delta_r-\Delta_s>\delta(\alpha_i,\beta_i)\quad\Rightarrow\quad f_i^*=s',.91,4.90,7.7,.75,color=ORANGE)
card(s,9.00,2.74,3.57,2.93,'经济含义','替换：成本低，但可能误伤合规者。\n整改：可能更贵，但保留关系、避免 Dα。',size=20)
strip(s,'相反的不等式选择整改；等号时无差异。治理阈值只依赖本供应商的审计。',y=6.50,size=17)

# 14 objective
s=slide('四、模型构建','买方的优化问题：最小化期望损失，同时满足总预算','原文 eqs. (9)–(12)，p. 10；期望损失为等价整理式',
        '建议用时 100 秒。买方最大化固定交易利润 Π 减两家供应商的期望损失。给定审计误差，p 是失败概率，整改总损失 pΔr 加漏报损失；替换总损失 pΔs 加误报声誉损失及漏报损失，取二者较小者。未审计损失为 (1−e)Dβ。基准中审计成本只进入预算约束，不从 Π 中重复扣除。代入最优治理后，问题有离散覆盖决策、连续误差决策和二次预算约束。i 为未审计时相关误差变量不影响收益。')
eq(s,r'\max\ \Pi-\sum_{i=1}^{2}\left[(1-x_i)(1-e)D_\beta+x_i\ell_i(\alpha_i,\beta_i)\right]',.84,2.48,11.55,.94)
eq(s,r'\ell_i=\min\{p_i\Delta_r,\ p_i\Delta_s+e\alpha_iD_\alpha\}+(1-e)\beta_iD_\beta',.88,3.70,11.48,.53)
eq(s,r'\sum_{i=1}^{2}\frac{x_i}{2}A_i^2\leq B,\qquad x_iA_i\geq x_ie',.94,4.47,7.0,.87)
eq(s,r'A_i=1-e\alpha_i-(1-e)\beta_i,\quad \alpha_i,\beta_i\in[0,1],\quad x_i\in\{0,1\}',.91,5.60,11.5,.57)
tx(s,8.00,4.62,4.57,.99,'pᵢ = eαᵢ + (1 − e)(1 − βᵢ)\n先代入最优治理，再求审计方案。',18,color=MUTED)
strip(s,'基准 B 是预算上限；第 6.2 节才将审计支出计入净利润。',y=6.55,size=17)

# 15 solution method
s=slide('四、模型构建','求解路径：后向归纳 + 内外层分解 + 覆盖数比较','§4–5，pp. 6–14；方法流程为汇报整理',
        '建议用时 65 秒。先解末阶段治理，得到后验阈值。再用一对一基准给出固定准确度的极端最优策略。扩展到两家时，先固定 b1、b2 解内层错误组合，再在预算圆弧上解外层分配。最后比较不审计、审计一家和两家三种覆盖方案。本章节沿用模板的“均衡分析”名称，但外生合规概率的主模型本质是买方的序贯最优决策，而不是供应商战略博弈。')
steps=[('治理最优响应','后验风险 → 整改 / 替换'),('一对一基准','固定 b → 选 α、β'),('两家内层','给定 b₁、b₂ → 错误政策'),('两家外层','预算约束 → 最优 b₁、b₂')]
for i,(a,b) in enumerate(steps):
    x=.72+i*3.10
    card(s,x,2.82,2.84,2.27,a,b,color=ORANGE if i==3 else BLUE,size=19)
    if i<3:arrow(s,x+2.88,3.72,.18,.22)
box(s,2.0,5.55,9.3,.84,fill=WARM)
tx(s,2.17,5.72,8.96,.47,'最后比较覆盖方案：不审计 vs 审计 1 家 vs 审计 2 家',22,color=ORANGE,bold=True,align=PP_ALIGN.CENTER)
tx(s,.92,6.60,11.52,.33,'主模型为单个买方的序贯优化；供应商内生努力拓展才引入策略互动。',17,color=MUTED)

# 16 extreme policies
s=slide('五、均衡分析','一对一基准：最优错误组合落在等准确度线的两个端点','§4；Lemma 2 / Proposition 1，pp. 7–8',
        '建议用时 75 秒。条件是买方决定审计，且预算能够支持有效审计。最大准确度 b0=min{1,√(2B)}。给定准确度时损失函数在整改和替换中取较小值，形成分段线性函数，最优误差可取端点。NFA 令 α=0，但只要 b<1 就允许 β>0；CA 令 β=0，但只要 b<1 就允许 α>0。若 e=1，公式 β 分母退化，需作为边界情形处理。审计门槛还要求治理成本低于违规损失，并满足最低资源要求。')
eq(s,r'b_0=\min\{1,\sqrt{2B}\}',.85,2.56,11.3,.65)
table(s,['策略','最优错误率','治理含义'],[
 ['No-false-alarm\n避免误报','α* = 0\nβ* = (1 − b₀)/(1 − e)','未通过者必为违规者；\n治理动作由成本比较决定。'],
 ['Catch-all\n全面筛查','α* = (1 − b₀)/e\nβ* = 0','不放过违规者；\n未通过者可能合规，治理需防误伤。']],y=3.59,rowh=.93,size=19,col=[3.39,4.19,4.35])
strip(s,'审计需要有治理价值：min{Δs, Δr} < Dβ；并满足最低准确度的预算要求。',y=6.6,size=17)

# 17 figure2
s=slide('五、均衡分析','原文图 2：治理成本决定优先控制哪类错误','原文 Figure 2 / Proposition 1，p. 8；e = 0.5，Dβ = 5.5，b₀ = 0.8',
        '建议用时 90 秒。横轴是整改成本，纵轴是替换成本；绿色横条表示 NFA，橙色竖条表示 CA，图中的 rectify/replace 又标记治理动作，二者不要混淆。左图 Dα=3，右图忽略一类损害 Dα=0。两类治理都很贵时，无效误判的治理浪费大，偏向避免误报；某一种治理很便宜时，偏向全面筛查。右上角两类治理成本都超过 Dβ 时不审计。中间政策与治理切换区才产生后续预算非单调结果。正文讲机制，完整不等式见附录。')
pic(s,ROOT/'fig2.png',.62,2.52,8.62,4.38);figures_used.add(2)
tx(s,9.46,2.71,3.04,.54,'读图顺序',21,color=BLUE,bold=True)
tx(s,9.46,3.41,3.02,2.92,'① 横轴 Δr，纵轴 Δs。\n② 纹理表示错误政策。\n③ rectify / replace 表示治理动作。\n④ 右图是 Dα = 0 对照。',18)
tx(s,9.46,6.05,3.00,.81,'高成本：优先防误报。\n低成本：优先防漏报。',16,color=ORANGE,bold=True)

# 18 budget shift numeric
s=slide('五、均衡分析','预算触发政策切换：企业损失下降，社会筛查却可能变弱','依据 Proposition 1(iii.b) / Proposition 2 自行计算；不是论文原图',
        '建议用时 95 秒。数值参数 e=.5,Dα=3,Dβ=5.5,Δs=1.6,Δr=2，位于 (Dβ−Dα)/2<Δs<Δr<Dβ/2 区间。条件为已审计单个供应商。低准确度采用 CA+整改，损失为 (2−e−b)Δr；高准确度 NFA+替换，损失为 (1−e)Δs+(1−b)(Dβ−Δs)。两条线相交于 b≈.8947，B≈.4003。b=.70 时选择 CA+整改，损失1.60且社会风险0；b=.95 时选择 NFA+替换，损失.995但社会风险.05。在阈值上，从0跃升到约.1053，之后随准确度提高下降。这里比较经济最优审计政策，不声称社会风险处处随预算增加。')
pic(s,ROOT/'benchmark_calc.png',.71,2.55,11.96,3.39)
tx(s,.84,6.03,11.71,.41,'算例：e = 0.5，Dα = 3，Dβ = 5.5，Δs = 1.6，Δr = 2；切换点 b̄̄ ≈ 0.895。',17,color=MUTED)
strip(s,'低预算：Catch-all + 整改；高预算：No-false-alarm + 低成本替换。',y=6.53,size=18)

# 19 ignore type I
s=slide('五、均衡分析','忽略误报损失：既可能选错错误政策，也可能选错治理动作','原文 Figure 2(b) / Corollary 1，pp. 8–9',
        '建议用时 65 秒。红框区域是原文标明忽略 Dα 后可能决策不同的区域。失误有两类：错误地选择 CA 使误报风险暴露，或错误地替换本应整改的供应商。特别强调“忽略 type I damage”是令 Dα=0，不是强制令 α=0，也不是消除误报本身。误报仍会造成治理浪费。论文中的损失比较是把错误决策放回真实 Dα>0 的环境评估，不是仅比较两个不同环境下的收益。')
pic(s,ROOT/'fig2b.png',.83,2.55,5.72,4.2)
card(s,6.83,2.72,5.52,1.44,'失误 1：政策选错','过度强调漏报，低估错误替换的声誉风险。',color=ORANGE,size=19)
card(s,6.83,4.35,5.52,1.66,'失误 2：动作选错','审计政策可能正确，但把本应整改的供应商替换，暴露于 Dα。',size=19)
tx(s,6.98,6.14,5.20,.77,'预算紧、Dα 大、Dβ 较小时，\n忽略误报的损失更明显。',17,color=BLUE,bold=True)

# 20 figure3
s=slide('五、均衡分析','原文图 3：给定不同预算，两个供应商可采用不同政策','原文 Figure 3 / Proposition 4，p. 12；e = 0.5，Dβ = 5.5，b₁ = 0.88，b₂ = 0.72',
        '建议用时 80 秒。图3给定 b1=.88,b2=.72，是内层错误政策问题；它本身没有证明这组预算是最优分配。左图考虑 Dα=3，右图 Dα=0。成本组合处于切换区时，准确度较高者可享受低成本替换，较低者因误判风险需要整改；某些区域也会产生 NFA/CA 的混合。随后命题5才求外层最优预算。图中 b 是最大可达准确度，而实际资源金额为 b²/2。')
pic(s,ROOT/'fig3.png',.62,2.53,8.68,4.34);figures_used.add(3)
tx(s,9.52,2.7,2.92,.46,'这一图回答内层问题',20,color=BLUE,bold=True)
tx(s,9.52,3.49,2.91,2.50,'b₁ > b₂：\n供应商 1 的审计更准。\n\n中间区域可能出现不同的治理动作，或不同的错误政策。',18)
tx(s,9.52,6.10,2.94,.80,'“给定不均分”不等于“已经证明不均分最优”。',17,color=ORANGE,bold=True)

# 21 uneven allocation
s=slide('五、均衡分析','外层预算分配：对称网络也可能“不均等配置”','Proposition 5 / Corollary 3，pp. 11–13；图为按模型自行计算',
        '建议用时 95 秒。此页条件是审计两家，图中采用前页算例 e=.5,Dα=3,Dβ=5.5,Δs=1.6,Δr=2。对每个 B 在二次预算可行圆弧上枚举8001个候选分配，并比较四类错误政策及治理损失。不同供应商可以位于不同政策分支，边际资源价值不同，因此对称不强制均分。低预算都整改，高预算都替换，中间可把一方推入低成本替换分支。此数值图是作者命题的解释性重算，不是原图或阈值的精确解析证明。命题5还允许两家都用CA但治理动作不同的不均分情形，不能把不均分限定为NFA/CA混合。')
pic(s,ROOT/'allocation_calc.png',.72,2.56,8.03,3.6)
card(s,9.01,2.79,3.60,2.90,'不均分的机制','同样的一元资源，\n在不同政策 / 治理分支上带来的损失降低不同。\n先让一家跨越切换点。',size=19)
tx(s,.88,6.16,11.69,.31,'算例：e = 0.5，Dα = 3，Dβ = 5.5，Δs = 1.6，Δr = 2；条件：审计两家；预算金额为 bᵢ²/2。',14.5,color=MUTED)
strip(s,'典型结构：低预算均分 → 中等预算不均分 → 高预算重新均分。',y=6.56,size=18)

# 22 social impact
s=slide('五、均衡分析','社会效果：审计两家时，预算增加可能使违规风险跳升','Proposition 6，p. 13；图为同一算例自行计算，条件：审计两家',
        '建议用时 85 秒。社会指标不是企业利润，而是至少一个违规投入品进入最终产品的概率。令单家漏出概率 qi=(1−e)βi，独立下 L=1−(1−q1)(1−q2)。两家CA时为0；一家NFA且另一家CA时等于1−b1；高预算均分且两家NFA时 b1=b2=√B，所以 L=1−B。在政策切换点，风险可能向上跳；每个固定分支内更高准确度仍改善筛查。图以已审计两家为条件，不能把这条线直接解释为所有预算下全局最优审计范围的社会效果。')
eq(s,r'L^*=1-\left[1-(1-e)\beta_1^*\right]\left[1-(1-e)\beta_2^*\right]',.8,2.55,11.70,.63)
pic(s,ROOT/'social_calc.png',.73,3.45,7.67,3.30)
tx(s,8.68,3.58,3.95,3.15,'两家 Catch-all：L* = 0\n混合政策：L* = 1 − b₁*\n两家 No-false-alarm 且均分：L* = 1 − B\n\n跳升来自政策改变，不是准确度下降。',18)

# 23 scope figure4
s=slide('五、均衡分析','原文图 4：审计覆盖数可能随预算出现“1 → 2 → 1 → 2”','图4，p. 14｜左：Dα=.306，Dβ=.72，Δr=.3，Δs=.6；右：Dα=2.5，Dβ=3.6，Δr=1.3，Δs=.5',
        '建议用时 100 秒。图横轴 e，纵轴 B。固定一个 e，沿竖直方向增加预算观察审计数量。左图通常由不审计到一家再到两家。右图在一些先验参数下，中间预算可从两家回到一家，再回到两家。不是所有e都如此。左图 Dα=.306,Dβ=.72,Δr=.3,Δs=.6；右图 Dα=2.5,Dβ=3.6,Δr=1.3,Δs=.5。虚线 B=e² 是两家最低可行预算，实线下界 B=e²/2 是一家最低可行预算。右图不能与前面的自行计算参数混为一谈。e=1 是边界无差异情形。')
pic(s,ROOT/'fig4.png',.69,2.57,11.95,3.93);figures_used.add(4)
strip(s,'读法：固定横轴 e，竖直向上提高 B；右图的部分区间中，审计数量反而减少。',y=6.58,size=17)

# 24 scope mechanism
s=slide('五、均衡分析','为什么更多预算可能只审一家？覆盖与精度的收益发生切换','Proposition 7 / Corollary 5，pp. 14–15',
        '建议用时 80 秒。低预算时集中资源提升一家准确度，中低预算足以审两家时用CA并整改。中高预算若集中一家可低误报地实施便宜替换，买方会宁愿承受另一家的固定违规风险，而不是分散资源后两家都只能昂贵整改。预算进一步升高，两家都能获得足够准确度，又恢复全覆盖。该解释既可适用于两家CA而治理切换，也可适用于错误政策同时切换，不能简单认定图4右图必须是从CA切NFA。未审计者的预期损失是 (1−e)Dβ。Dα=0 时这类覆盖反转消失。')
table(s,['预算阶段','资源选择','经济机制'],[
 ['低','集中审计 1 家','先获得足够的有效准确度'],
 ['中低','审计 2 家','覆盖收益提升；失败后偏向整改'],
 ['中高（特定参数）','重新集中审计 1 家','一家的高精度允许低成本替换；\n接受另一家的固定预期风险'],
 ['高','审计 2 家','两家均可降低错误，重新获得覆盖优势']],rowh=.78,size=18,col=[2.52,3.51,5.90])
strip(s,'Dα = 0 时，审计覆盖随预算扩张；误报损失使这一直觉不再普遍成立。',y=6.65,size=17)

# 25 synthesis
s=slide('五、均衡分析','三条反直觉结果，来自同一个联动机制','综合 Proposition 1、5–7；所有结论均有参数条件',
        '建议用时 50 秒。这一页收束分析：根源是成本优势与误报风险共同导致政策和治理动作切换。预算仅扩大可行集合，企业最优收益弱增加，但社会指标不必单调。三个结果都是“可能”，不是一般情况下必然不均分、必须减少覆盖或预算一定有害。')
card(s,.72,2.72,3.84,2.55,'对称 ≠ 必须均分','政策 / 治理分支不同，资源的边际价值不同。',size=21)
card(s,4.75,2.72,3.84,2.55,'预算多 ≠ 覆盖更多','高精度带来的低成本治理优势，可能超过多审一家的收益。',color=ORANGE,size=20)
card(s,8.78,2.72,3.84,2.55,'私益提升 ≠ 风险下降','从 Catch-all 转向 No-false-alarm，可改善收益却允许漏报。',size=20)
box(s,1.10,5.78,11.14,.82,fill=WARM)
tx(s,1.26,5.96,10.82,.44,'预算 → 准确度 → 错误政策 → 治理动作 → 资源配置与社会结果',20,color=ORANGE,bold=True,align=PP_ALIGN.CENTER)

# 26 effort extension
s=slide('六、扩展模型','拓展 1：供应商主动选择责任努力，审计会反向塑造激励','§6.1 / Proposition 8，pp. 15–16；一对一渠道拓展',
        '建议用时 95 秒。该拓展针对一对一基准，不是完整两供应商装配博弈。供应商先选择 eS，成本 κeS²/2，然后买方观察努力选择审计及治理。供应商在交易中获得 ΠS，被整改自己承担 δr，被替换利润为0。δr 与买方整改成本 Δr 不同。命题8假定 κ/2<δr<ΠS<κ 及足够预算。稳定整改分支努力 erS=δr/κ，稳定替换分支 esS=ΠS/κ，后者更高。中间区间存在边界努力，使买方正好愿意保留供应商，完整分段结果见附录，不能只用两个候选值覆盖所有情况。')
card(s,.73,2.66,3.82,1.60,'供应商先行动','选择 eS；成本 κeS²/2。\n买方观察 eS 后设计审计。',size=19)
arrow(s,4.68,3.31,.36,.28)
card(s,5.20,2.66,3.33,1.60,'治理影响收益','整改：供应商付 δr。\n替换：供应商利润为 0。',color=ORANGE,size=18)
arrow(s,8.65,3.31,.36,.28)
card(s,9.18,2.66,3.43,1.60,'供应商预期响应','保住交易关系的价值，决定主动合规的动力。',size=19)
eq(s,r'e_S^r=\frac{\delta_r}{\kappa}\ <\ e_S^s=\frac{\Pi_S}{\kappa}',1.06,4.71,11.06,.77)
tx(s,.94,5.68,11.46,.64,'稳定整改 / 替换分支：替换威胁提高合规努力；中间区间还可能出现阈值努力。',21)
tx(s,.94,6.57,11.42,.35,'前提：κ/2 < δr < ΠS < κ；并满足命题 8 的有效准确度与预算条件。',16,color=MUTED)

# 27 random extension
s=slide('六、扩展模型','低预算下的随机审计：通过事前承诺协调努力与审计能力','§6.1 / Proposition 9，pp. 16–17',
        '建议用时 65 秒。如果纯策略候选合规努力 ePS 超过买方最大准确度 b0，就违反了有效审计条件；一旦不审计，供应商可能降低努力，买方又想审计，导致纯策略候选难以维持。文章提出买方事前承诺以 q 的概率启动审计。命题9在 1/2<b0<ePS 且 ePS属于稳定整改或替换分支时，用 q=b0/ePS 支持一个均衡。不要把这描述为证明所有低预算环境都有同一个混合策略纳什均衡；它是特定条件下的承诺随机政策。')
card(s,.83,2.77,5.60,2.82,'纯策略可能无法维持','供应商努力高于可达审计准确度\n→ 有效审计条件不满足\n→ 买方停止审计\n→ 供应商又降低努力。',size=20)
card(s,6.74,2.77,5.60,2.82,'事前随机承诺','买方以 q ∈ (0, 1) 的概率审计。\n降低供应商的有效审计暴露，\n使诱导努力与预算能力一致。',color=ORANGE,size=20)
eq(s,r'\frac{1}{2}<b_0<e_S^P,\qquad q=\frac{b_0}{e_S^P}',1.04,5.97,11.23,.69)

# 28 endogenous budget
s=slide('六、扩展模型','拓展 2：预算内生化，以治理收益与审计边际成本确定投入','§6.2 / Proposition 10，p. 17；条件：一对一且审计最优',
        '建议用时 95 秒。恢复成本系数 c，买方现在最大化 π*(B)−B，因此不仅在预算内分配，还要决定是否值得投入。每个固定分支的期望毛利对 A 是线性的，斜率 w 为提高准确度的边际收益；减去 cA²/2 后一阶候选 A*=w/c，B*=w²/(2c)。表中列4个分支的 w，最终要比较分支是否自洽、A是否在(e,1]，以及不审计的收益。不能不加条件地使用 w/c 超过1的结果。核心切换区，c较低时偏向 NFA+替换，较高时偏向CA+整改，审计技术效率影响的不只是准确度。')
eq(s,r'\max_B\{\pi^*(B)-B\},\qquad B=\frac{c}{2}A^2',.89,2.56,11.6,.61)
table(s,['候选政策 / 治理','准确度边际收益 w','内点候选'],[
 ['No-false-alarm + 整改','Dβ − Δr','A* = w/c'],
 ['Catch-all + 整改','Δr','B* = c(A*)²/2'],
 ['No-false-alarm + 替换','Dβ − Δs','比较候选分支与不审计'],
 ['Catch-all + 替换','Dα + Δs','检查 A* ∈ (e, 1] 及分支条件']],y=3.50,rowh=.54,size=17,col=[4.04,3.45,4.44])
strip(s,'技术效率提高（c 下降）可能改变治理模式，而不只是提高同一模式下的准确度。',y=6.52,size=17)

# 29 conclusions
s=slide('七、结论与未来展望','主要结论：责任审计应与治理动作、资源配置共同设计','总结 §7，p. 18；结合 Proposition 1–10',
        '建议用时 65 秒。把结论带回最初四个研究问题。第一，不能忽略误报，因为它不仅改变声誉损失，还会改变治理动作。第二，政策依赖成本和可达准确度，不能只问审计力度多大。第三，中等预算下可能出现不均分和覆盖反转。第四，企业利润最优与社会筛查最优有差距。拓展说明供应商激励与技术成本又会反馈至政策。')
bullets(s,[
 '两类错误必须联合管理：误报会改变治理方式，也会改变预算与覆盖决策。',
 '成本与预算共同决定政策：有限资源下可能选择全面筛查 + 整改。',
 '资源配置可能非单调：对称不保证均分；更多预算不保证覆盖更多。',
 '企业收益与社会筛查需分别评估：控制误报可能伴随正的漏报风险。'],gap=.91,size=22)

# 30 appraisal
s=slide('七、结论与未来展望','模型评价：结论清楚，但依赖若干关键结构假设','此页为汇报者评述；依据 §3–6 的模型设定',
        '建议用时 65 秒。该页明确是汇报者评价，不是作者已验证的结果。线性等准确度线和自由选择误差组合使端点政策有明确表达；真实审计可行集合可能是受技术限制的ROC曲线。整改和替换完全成功使漏报是主要残余风险，若治理会失败则结论可能改变。两家对称、状态独立、逐供应商损失可加总使治理可分解，只有共享预算耦合；相关违规状态或非线性品牌声誉损失会改变分解。基准供应商不策略性响应，第6.1节仅在一对一扩展了努力内生化，不宜视为完整网络稳健性检验。')
table(s,['关键假设','作用','可讨论的现实差异'],[
 ['等准确度线 + 二次成本','得到清晰的端点政策与预算分配','真实误差组合受审计技术 / ROC 限制'],
 ['整改 / 替换后完全合规','失败被治理后不再形成残余违规','整改效果、替换质量可能不确定'],
 ['对称、独立；损失可加总','治理逐家可分；预算形成耦合','供应商风险相关；品牌损失可能非线性'],
 ['基准 e 外生；单期决策','隔离两类错误与治理的联动','努力激励、周期审计和学习会反馈']],rowh=.79,size=16.5,col=[3.79,4.0,4.14])
tx(s,.89,6.69,11.61,.3,'这些是模型适用边界与研究机会，不能据此直接判定论文结论不成立。',16,color=MUTED)

# 31 future
s=slide('七、结论与未来展望','未来方向：从单买方、单期和简单装配，走向网络协调','作者展望：§7，p. 18；末项明确为汇报者补充',
        '建议用时 45 秒。前三项来自作者结论部分：多个竞争买方共享上游供应商时存在外部性，需要协调；多层网络可能直接审计或委托审计，涉及激励；周期审计可与采购和风险状态动态演化结合。最后一项是汇报者补充：把真实ROC可行集合、治理失败及相关状态纳入，检验端点策略和非单调性是否稳健。区分作者提出的方向和个人建议。')
bullets(s,[
 '多买方与共享供应商：竞争、责任风险外部性与合作审计机制。',
 '多层级供应网络：直接审计、委托审计及上游监督激励。',
 '周期审计与采购联动：动态风险、信息学习和长期治理安排。',
 '汇报者补充：受技术限制的错误率组合、治理失败与相关合规状态。'],gap=.93,size=21)

# 32 Thanks
s=slide('讨论与交流','',notes='正文到此结束。可讨论三个问题：如果整改可能失败，CA是否仍然保证零违规产品？如果误报只影响供应商、不影响买方声誉，谁会承担校验成本？如何用现实审计数据识别α、β与治理成本以检验覆盖反转？附录为讨论时备用。')
tx(s,2.02,2.56,9.28,1.70,'感谢聆听\n敬请批评指正',37,color=BLUE,bold=True,align=PP_ALIGN.CENTER)
tx(s,1.20,4.88,10.94,.94,'讨论：企业最优审计方案，应如何兼顾供应商权益与社会责任风险？',22,align=PP_ALIGN.CENTER)
tx(s,2.09,6.21,9.14,.48,'附录：命题条件、损失函数与拓展结果',18,color=MUTED,align=PP_ALIGN.CENTER)

# 33 appendix exact one-channel
s=slide('附录 A｜一对一完整政策条件','Proposition 1：审计政策的分区条件','原文 Proposition 1，pp. 7–8；省略等号无差异情形',
        '附录供讨论查阅，不在30分钟主线中逐条念。条件为买方决定审计且可达准确度有效。这里将原文命题1完整主要分区译为表格。若Δs>Δr，整改便宜，政策由Δr与Dβ/2决定。若Δr>Δs且超出中间区间，政策由Δs与(Dβ−Dα)/2决定。中间区间中，若替换成本极低就总是CA，但治理仍能随预算切换；若替换成本在中等区间，政策会在b0跨过bbar后由CA转NFA。不能把极低替换成本区的不均预算误解为两家用了不同错误政策。',True)
table(s,['参数区间','额外条件','最优错误政策'],[
 ['Δs > Δr','Δr > Dβ/2；反之','NFA；反之 CA'],
 ['Δr > Δs；\nΔr > min{Δs + Dα/2, Dβ/2}','Δs > (Dβ − Dα)/2；反之','NFA；反之 CA'],
 ['Δs < Δr < min{Δs + Dα/2, Dβ/2}','Δs < (Dβ − Dα)/2','始终 CA'],
 ['同上','Δs > (Dβ − Dα)/2；\nb₀ > b̄̄；反之','NFA；反之 CA']],rowh=.75,size=16.5,col=[4.95,4.05,2.93])
eq(s,r'\overline{\overline{b}}=\frac{D_\beta-(2-e)\Delta_r-e\Delta_s}{D_\beta-\Delta_r-\Delta_s}',.95,6.36,11.32,.57)

# 34 appendix loss affine
s=slide('附录 B｜四种损失分支','固定准确度 b 时，比较四条线性期望损失','按 §3–4 等价推导；本文算例与预算数值图的计算基础',
        '这四个式子由事件表逐项累加得到。NFA中失败概率是 b−e，CA中失败概率是2−e−b。NFA无误报损害，CA误报损害的无条件概率为1−b。每个b取四项的最小值，同时未审计损失为(1−e)Dβ。数值图对称下约束b1≥b2，沿两家预算圆弧枚举所有候选分配；每个候选比较全部四种政策，避免预先锁定治理动作。此表是汇报推导而不是复制原文新命题。',True)
table(s,['组合','每家期望损失 ℓ(b)'],[
 ['NFA + 整改','(b − e)Δr + (1 − b)Dβ'],
 ['NFA + 替换','(b − e)Δs + (1 − b)Dβ'],
 ['CA + 整改','(2 − e − b)Δr'],
 ['CA + 替换','(2 − e − b)Δs + (1 − b)Dα'],
 ['不审计','(1 − e)Dβ']],rowh=.62,size=20,col=[4.28,7.65])
strip(s,'ℓ*(b) 取四个分支的最小值；π* = Π − ℓ*（两家时对期望损失求和）。',y=6.53,size=17)

# 35 appendix allocation conditions
s=slide('附录 C｜预算分配的条件','不均分不只出现在 NFA / CA 混合政策下','Proposition 5 / Corollary 3，pp. 11–13',
        '命题5条件一：Δs>Δr，或者Δr>Δs且Δr>min{Δs+Dα/2,Dβ/2}，预算总是均分。条件二的两种子情况为(a)Δs<Δr<Δs+Dα/2且Δs<(Dβ−Dα)/2，(b)(Dβ−Dα)/2<Δs<Δr<Dβ/2，均可在中等预算出现不均分。其中a两家仍然都采用CA，只是整改与替换动作可能不同；b还可能出现一方NFA另一方CA。这区分是汇报中容易混淆的关键。完整BL与BU闭式定义在补充材料，此上传正文没有EC附件，因此不伪造阈值表达式。',True)
table(s,['分配结构','参数条件'],[
 ['始终均分','Δs > Δr；或 Δr > Δs 且\nΔr > min{Δs + Dα/2, Dβ/2}'],
 ['中等预算不均分：情形 (a)','Δs < Δr < Δs + Dα/2，且\nΔs < (Dβ − Dα)/2'],
 ['中等预算不均分：情形 (b)','(Dβ − Dα)/2 < Δs < Δr < Dβ/2']],rowh=.88,size=18,col=[4.23,7.70])
tx(s,.88,6.08,11.59,.84,'情形 (a)：两家都可用 CA，治理动作不同。\n情形 (b)：中等预算可用 NFA / CA 混合；高预算两家 NFA。',17,color=BLUE,bold=True)

# 36 appendix social indicator audit scope
s=slide('附录 D｜社会指标与覆盖范围','社会风险要连同审计覆盖一起解释','Proposition 2、6；覆盖一般式为基于模型的汇报推导',
        '该公式把未审计与被审计统一起来。在独立供应商状态下，单家残余违规概率是 qi=(1−xi)(1−e)+xi(1−e)βi，治理成功使失败分支没有残余风险。最终装配产品只要任一输入违规就违规，所以总体是1−乘积。未审计两家为1−e²。只审一家CA也仍有未审计者风险1−e，不能宣称整个网络零违规。只审一家NFA时为1−e b0。主文命题6以审计两家为条件，而本文算例的社会曲线也使用该条件。',True)
eq(s,r'q_i=(1-x_i)(1-e)+x_i(1-e)\beta_i',.95,2.60,11.29,.68)
eq(s,r'L=1-(1-q_1)(1-q_2)',.96,3.55,11.24,.61)
table(s,['覆盖 / 政策','最终产品不合规概率'],[
 ['不审计任何一家','1 − e²'],
 ['审计一家且用 CA','1 − e'],
 ['审计一家且用 NFA','1 − e b₀'],
 ['审计两家且均用 CA','0']],y=4.43,rowh=.48,size=17,col=[6.02,5.91])

# 37 appendix prop8 original
s=slide('附录 E｜供应商内生努力','Proposition 8：完整分段结果与适用条件','原文 Proposition 8，pp. 15–16；左侧保留原文命题片段',
        '左侧为论文第16页命题8的主要内容。其假设κ/2<δr<ΠS<κ在第15页命题开头，见本页右侧。复杂中间区间用边界努力e-bar/e-doublebar维持买方对治理或错误政策的无差异。ε是足够小的扰动；阈值ΠS-bar等在原命题定义及补充证明中，不在汇报中擅自简化。这是单供应商拓展，不是两家装配网络预算分配的完整内生努力求解。',True)
pic(s,ROOT/'prop8.png',.80,2.50,5.18,4.46)
card(s,6.42,2.81,5.74,1.59,'适用前提','κ/2 < δr < ΠS < κ；\nb₀ ≥ max{eS*, 1 − eS*} ≥ 1/2。',size=19)
card(s,6.42,4.69,5.74,1.84,'读法','先辨认稳定整改 / 稳定替换分支，\n再看切换区中的边界努力。\nδr 与买方成本 Δr 是不同参数。',color=ORANGE,size=19)

# 38 appendix prop10 original
s=slide('附录 F｜内生预算的分区','Proposition 10：分支候选需要与边界和不审计方案比较','原文 eqs. (18)–(21) / Proposition 10，p. 17',
        '左侧为原文命题10，右上保留内生预算优化约束。原文讨论的是审计最优且候选准确度在(e,1]的非平凡区间；超出该条件时，要检查完全准确边界、最低有效准确度边界或不审计。一个分支内若w为常量，A*=w/c、B*=w²/(2c)，因此成本系数上升使投入减少；跨分支时要比较净收益。不存在任意c下直接把未裁剪w/c当作完整全局解的结论。',True)
pic(s,ROOT/'prop10.png',.81,2.56,5.61,4.35)
pic(s,ROOT/'eq18.png',7.20,2.67,4.37,1.84)
card(s,6.78,4.83,5.59,1.64,'使用命题时的检查','① 审计是否最优；② 分支是否自洽；\n③ e < A* ≤ 1；④ 净利润比较。',color=ORANGE,size=18)

timing=[20,25,15,55,65,45,45,55,55,60,55,55,85,90,55,65,70,80,50,65,80,75,85,65,45,80,55,85,45,45,30]
assert sum(timing)==1800
for i,seconds in enumerate(timing):
    notes_all[i]['notes']=re.sub(r'建议用时 \d+ 秒',f'建议用时 {seconds} 秒',notes_all[i]['notes'])
    prs.slides[i].notes_slide.notes_text_frame.text=notes_all[i]['notes']
outfile=OUT/'deck.pptx'
prs.save(outfile)
(ROOT/'slide_notes.json').write_text(json.dumps(notes_all,ensure_ascii=False,indent=2))
md=['# 组会汇报讲解备注','正文 32 页，建议约 30 分钟；附录 6 页供讨论。',
    '论文：Li, C. & Li, Y. (2026), Catch-all or no-false-alarm? Managing socially responsible supply chains with imperfect audits. Production and Operations Management. DOI: 10.1177/10591478261468995.',
    '模板：用户提供的《大组会.pptx》；沿用原背景、南京审计大学校徽、蓝橙配色及七部分结构。',
    '封面汇报人与汇报日期可在 PowerPoint 内填写。']
for item in notes_all:
    md.extend([f"\n## 第 {item['slide']} 页 · {item['title'] or item['section']}",item['notes']])
(OUT/'组会汇报_讲解备注.md').write_text('\n\n'.join(md))
print(json.dumps({'pptx':str(outfile),'slides':len(prs.slides),'main_slides':32,'appendix_slides':6,'original_figures':sorted(figures_used),'equation_images':math_count,'policy_threshold':bbar},ensure_ascii=False))
